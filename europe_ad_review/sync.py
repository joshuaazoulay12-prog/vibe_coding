import openpyxl, pickle, copy, datetime
rows, ADDED, CH, idmap, N = pickle.load(open('improve.pkl','rb'))
wb = openpyxl.load_workbook('Europe_Improved.xlsx')
ws = wb['Europe']
H = [ws.cell(1,c).value for c in range(1,29)]; IDX={h:i for i,h in enumerate(H)}
def g(x,k): return x[IDX[k]]

# ---------- 1. Master_Deals Europe block ----------
md = wb['Master_Deals']
START = 91
old_last = 219
# clear old block
for r in range(START, old_last+1):
    for c in range(1, 26): md.cell(r,c).value = None
tf, ta, tb = md.cell(92,1).font, md.cell(92,1).alignment, md.cell(92,1).border
need = len(rows) - (old_last-START+1)
if need > 0: md.insert_rows(old_last+1, need)
for i, x in enumerate(rows):
    r = START + i
    vals = [g(x,'Deal ID'),'Europe',g(x,'Target Country'),g(x,'Target / Business'),g(x,'Acquirer'),
            g(x,'Acquirer Country'),g(x,'Seller'),g(x,'Announcement Date'),g(x,'Completion / IPO Date'),
            None, g(x,'Transaction Type'), g(x,'Deal Structure'), g(x,'Deal Status'),
            g(x,'Deal Value (USD mn)'), g(x,'Currency'), g(x,'Hardware / Software / Services'),
            g(x,'A&D Domain'), g(x,'Inclusion Tier'), g(x,'Scope Evidence'), g(x,'Verification Status'),
            g(x,'Primary Source URL'), g(x,'Verification Status'), g(x,'Audit Notes')]
    d = g(x,'Completion / IPO Date')
    vals[9] = d.year if isinstance(d, datetime.datetime) else None
    for c, v in enumerate(vals, 1):
        cell = md.cell(r,c); cell.value = v
        cell.font=copy.copy(tf); cell.alignment=copy.copy(ta); cell.border=copy.copy(tb)
    md.cell(r,9).number_format='yyyy-mm-dd'; md.cell(r,8).number_format='yyyy-mm-dd'
print(f"Master_Deals Europe block: rows {START}-{START+len(rows)-1} ({len(rows)} rows)")

# ---------- 2. Europe_Audit_Log ----------
al = wb['Europe_Audit_Log']
tf, ta, tb = al.cell(2,1).font, al.cell(2,1).alignment, al.cell(2,1).border
# remap Final Deal ID for retained rows
remap=0
for r in range(2, al.max_row+1):
    fid = al.cell(r,4).value
    if fid in idmap:
        al.cell(r,4).value = idmap[fid]; remap+=1
start = al.max_row+1
LOG = []
for tgt, note, src, url, did in [
 ('Milrem Robotics','ADDED under the strict policy: Estonian target, defence robotics as sole business, EDGE Group acquired a majority stake announced 15 February 2023. Absent from the original workbook. No separate legal completion date was disclosed by either party.','EDGE Group announcement (IDEX 2023)','https://edgegroup.ae/news/edge-acquires-majority-stake-milrem-robotics-europes-leading-developer-robotics-and-autonomous', None),
 ('CS Group','ADDED under the strict policy: French defence and security critical-systems prime; Sopra Steria finalised a 75.06% majority on 28 February 2023. Absent from the original workbook, although CS Group appears in the sheet as an acquirer.','Sopra Steria completion press release, 28 February 2023','https://www.soprasteria.com/newsroom/press-releases/details/sopra-steria-finalises-its-acquisition-of-a-majority-stake-in-the-share-capital-of-cs-group', None),
 ('TKMS AG & Co. KGaA','ADDED: admitted to trading in the Prime Standard of the Frankfurt Stock Exchange on 20 October 2025, the largest European pure-play naval defence listing of the review period (~EUR 6.8bn market value at debut, ~EUR 18.6bn order backlog). SCOPE CAVEAT: this is a spin-off admission with no public offering and no proceeds, so it does not satisfy the workbook\'s strict "Completed IPO" definition. Flagged in Deal Structure so it can be filtered.','TKMS Group listing release','https://www.tkmsgroup.com/news/article/tkms-lists-on-the-frankfurt-stock-exchange-successful-market-debut-for-maritime-defense-provider', None),
 ('Spirit AeroSystems Belfast and Prestwick operations','ADDED under the strict policy: UK aerostructures sites (A220 wings and mid-fuselage; A320/A350 wing components) carved out to Airbus, completed 8 December 2025, over 4,000 employees transferred. Absent from the original workbook. Only the European sites are recorded; Kinston, NC is out of scope.','Airbus completion press release, December 2025','https://www.airbus.com/en/newsroom/press-releases/2025-12-airbus-completes-acquisition-of-spirit-aerosystems-sites', None),
]:
    LOG.append(('n/a — not in original','n/a', tgt, None, tgt, 'add', 'verified', note, 'See Europe sheet Scope Evidence', src, url))
for did, tgt, what, evid, url in CH:
    LOG.append(('see original mapping', did, tgt, idmap.get(did,did), tgt, 'correct', 'verified', what+' — '+evid, 'Scope unchanged', 'Primary source (see URL)', url))
for i, rec in enumerate(LOG):
    r = start + i
    for c, v in enumerate(rec, 1):
        cell = al.cell(r,c); cell.value = v
        cell.font=copy.copy(tf); cell.alignment=copy.copy(ta); cell.border=copy.copy(tb)
# fill Final Deal ID for additions
for r in range(start, start+4):
    t = al.cell(r,5).value
    for x in rows:
        if g(x,'Target / Business') == t: al.cell(r,4).value = g(x,'Deal ID')
print(f"Europe_Audit_Log: remapped {remap} final IDs, appended {len(LOG)} entries")

# ---------- 3. Sources_Audit ----------
sa = wb['Sources_Audit']
tf, ta, tb = sa.cell(2,1).font, sa.cell(2,1).alignment, sa.cell(2,1).border
for r in range(2, sa.max_row+1):
    if sa.cell(r,1).value in idmap: sa.cell(r,1).value = idmap[sa.cell(r,1).value]
s2 = sa.max_row+1
SRC=[]
for x in ADDED:
    SRC.append((g(x,'Deal ID'), g(x,'Target / Business'), 'Completion / IPO / scope',
                'Issuer or acquirer primary release', g(x,'Primary Source URL'), g(x,'Secondary Source URL'),
                '2026-08-02', 'Added in the 2026-08-02 review; absent from the original workbook and audit log.'))
for did, tgt, what, evid, url in CH:
    SRC.append((idmap.get(did,did), tgt, 'Corrected value or date', 'Issuer or acquirer primary release',
                url, 'See Europe sheet Secondary Source URL', '2026-08-02', what+' — '+evid))
for i, rec in enumerate(SRC):
    r = s2+i
    for c, v in enumerate(rec,1):
        cell = sa.cell(r,c); cell.value=v
        cell.font=copy.copy(tf); cell.alignment=copy.copy(ta); cell.border=copy.copy(tb)
print(f"Sources_Audit: appended {len(SRC)} entries")

# ---------- 4. Dashboard + Checks ranges ----------
LAST = len(rows)+1
def fix(sheet):
    n=0
    for row in sheet.iter_rows():
        for c in row:
            if isinstance(c.value,str) and c.value.startswith('=') and ('130' in c.value or 'A2:A130' in c.value):
                c.value = c.value.replace('130', str(LAST)); n+=1
    return n
print("Dashboard formulas updated:", fix(wb['Dashboard']))
print("Checks formulas updated:", fix(wb['Checks']))
ch = wb['Checks']
for r in range(2, ch.max_row+1):
    if ch.cell(r,1).value and 'Master_Deals Europe count' in str(ch.cell(r,1).value):
        ch.cell(r,2).value = len(rows); ch.cell(r,3).value = len(rows)
wb.save('Europe_Improved.xlsx')
print("saved")
