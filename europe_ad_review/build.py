import openpyxl, re, datetime
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
exec(open('verif.py').read())
exec(open('corrections.py').read())
exec(open('additions.py').read())

src = openpyxl.load_workbook('original.xlsx', data_only=False)
ws  = src['Europe']
H   = {ws.cell(1,c).value: c for c in range(1, 37) if ws.cell(1,c).value}
def g(r, name): return ws.cell(r, H[name]).value

FX = {'EUR':1.09,'GBP':1.27,'SEK':0.095,'NOK':0.094,'DKK':0.146,'CHF':1.12,'HUF':0.0027,
      'PLN':0.25,'CZK':0.044,'INR':0.012,'USD':1.0,'JPY':0.0068,'CAD':0.735,'AUD':0.66}

def parse_local(txt):
    """Pull a USD-mn figure out of a Block-B 'Not disclosed in USD (...)' string."""
    if not isinstance(txt,str): return None, txt
    m = re.search(r'(EUR|GBP|SEK|NOK|DKK|CHF|HUF|PLN|CZK|USD)\s*([\d,]+(?:\.\d+)?)\s*(bn|billion|m\b|mn)?', txt, re.I)
    if not m: return None, txt
    cur=m.group(1).upper(); num=float(m.group(2).replace(',',''))
    if (m.group(3) or '').lower() in ('bn','billion'): num*=1000
    return round(num*FX.get(cur,1.0),2), txt

recs=[]; removed=[]; changelog=[]
for r in range(2, ws.max_row+1):
    did = str(g(r,'Deal ID'))
    if did in REMOVE:
        removed.append((did, g(r,'Target / Asset'), g(r,'Target Country'), g(r,'Acquirer'),
                        str(g(r,'Deal Date'))[:10], g(r,'Deal Value (USD mn)'), REMOVE[did]))
        continue
    c = C.get(did, {})
    dt = g(r,'Deal Date')
    orig_date = dt.strftime('%Y-%m-%d') if hasattr(dt,'year') else str(dt)
    ann  = c.get('ann')
    comp = c.get('comp', orig_date if hasattr(dt,'year') else None)
    if 'comp' in c and c['comp'] is None: comp = None
    status = c.get('stat', g(r,'Deal Status'))
    raw_val = g(r,'Deal Value (USD mn)')
    if isinstance(raw_val,str):
        val, local_txt = parse_local(raw_val)
        disc = 'Disclosed in local currency only' if val else 'Not disclosed'
        basis = 'Reported consideration (converted)' if val else 'n/a'
    else:
        val, local_txt, disc, basis = raw_val, None, 'Reported (vendor data)', 'Reported deal value'
    if 'val' in c: val = round(c['val'],2) if c['val'] is not None else None
    local_txt = c.get('loc', local_txt)
    disc      = c.get('disc', disc)
    basis     = c.get('basis', basis)
    n = int(re.sub(r'\D','',did))
    block = 'A (PitchBook export)' if n<=198 else 'B (primary-source research)'
    vd = V.get(did, (None,None,None))
    if vd[0]:
        reliab = {'OK':'Primary-source verified','FIX':'Primary-source verified - corrected',
                  'EST':'Primary-source verified - value undisclosed','SCOPE':'Primary-source verified'}.get(vd[0].split('/')[0],'Primary-source verified')
    elif n>=199: reliab = 'Vendor/primary source cited, not re-verified in this pass'
    else:        reliab = 'NOT INDEPENDENTLY VERIFIED - vendor data only'
    evt = 'IPO' if g(r,'Transaction Type')=='IPO' else 'Merger/Acquisition'
    tier = c.get('tier', g(r,'Inclusion Tier'))
    tier = {'Core / primarily defence':'Core A&D','Core / diversified A&D':'Core A&D'}.get(tier,tier)
    dom  = c.get('dom',  g(r,'A&D Domain'))
    scope = c.get('scope','Pass')
    src_url = c.get('src', vd[2] if vd[2] and vd[2].startswith('http') else g(r,'Source URL'))
    year = c.get('year', g(r,'Deal Year'))
    if c:
        changelog.append((did, g(r,'Target / Asset'), block,
                          '; '.join(f"{k}" for k in c if k!='note'),
                          c.get('note') or (vd[1] or ''), src_url or ''))
    recs.append(dict(legacy=did, tgt=c.get('tgt', g(r,'Target / Asset')), tc=c.get('tc', g(r,'Target Country')),
        hq=c.get('hq', g(r,'HQ Location')), acq=(None if 'acq' in c else g(r,'Acquirer')),
        ac=c.get('ac', g(r,'Acquirer Country')), seller=c.get('seller', g(r,'Seller')),
        evt=evt, struct=c.get('struct', g(r,'Deal Structure')), stake=c.get('stake','100%' if evt!='IPO' else 'n/a'),
        ann=ann, comp=comp, year=year, status=status, val=val, basis=basis, local=local_txt,
        disc=disc, tier=tier, dom=dom, scope=scope, desc=g(r,'Description'),
        syn=g(r,'Deal Synopsis'), reliab=reliab, src=src_url, note=c.get('note') or (vd[1] if vd[1] else None), block=block))

for a in ADD:
    (tgt,tc,hq,acq,ac,seller,evt,struct,stake,ann,comp,status,val,local,basis,disc,tier,dom,desc,url) = a
    recs.append(dict(legacy='(added)', tgt=tgt, tc=tc, hq=hq, acq=acq, ac=ac, seller=seller, evt=evt,
        struct=struct, stake=stake, ann=ann, comp=comp, year=int((comp or ann)[:4]), status=status,
        val=round(val,2) if val else None, basis=basis, local=local, disc=disc, tier=tier, dom=dom,
        scope='Pass', desc=desc, syn=None, reliab='Primary-source verified - ADDED', src=url,
        note='Verified in-scope event absent from the original Europe tab.', block='C (added in this pass)'))

recs.sort(key=lambda x: (x['comp'] or x['ann'] or '9999'))
for i,x in enumerate(recs,1): x['id'] = f"EU-{i:03d}"
print(f"corrected rows: {len(recs)} | removed: {len(removed)} | changelog entries: {len(changelog)}")
import pickle; pickle.dump((recs,removed,changelog), open('built.pkl','wb'))
